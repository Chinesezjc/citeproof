"""HTTP API for CiteProof.

An audit is submitted as a job and polled for progress, because a full audit
performs one or two upstream requests per citation and the anonymous access mode
paces those requests to stay inside the corpus allowance. Audits run on a single
worker so that several concurrent submissions cannot exceed that allowance.
"""

from __future__ import annotations

import io
import json
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles

from . import __version__
from .config import REPO_ROOT, Settings, load_settings
from .courtlistener import CourtListenerClient, ResponseCache
from .extract import extract_citations
from .llm import LLMClient
from .report import to_json, to_markdown
from .schemas import AuditRequest, AuditSummary, Verdict
from .verify import CitationVerifier, VerifierOptions

# The demonstration documents ship with the application so that a reviewer can
# run an audit immediately, without a document of their own.
EXAMPLES_DIR = REPO_ROOT / "data" / "examples"
AUDITS_DIR = REPO_ROOT / "data" / "audits"
BENCHMARK_RESULTS = REPO_ROOT / "data" / "benchmark" / "results.json"
MAX_DOCUMENT_CHARS = 400_000


class AuditJobs:
    """In-memory job registry with reports persisted to disk."""

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._jobs: dict[str, AuditSummary] = {}
        self._executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="audit")

    def create(self) -> AuditSummary:
        job = AuditSummary(audit_id=uuid.uuid4().hex[:12], status="queued", message="Queued")
        with self._lock:
            self._jobs[job.audit_id] = job
        return job

    def get(self, audit_id: str) -> AuditSummary | None:
        with self._lock:
            job = self._jobs.get(audit_id)
        if job is not None:
            return job
        path = AUDITS_DIR / f"{audit_id}.json"
        if path.is_file():
            payload = json.loads(path.read_text(encoding="utf-8"))
            return AuditSummary(
                audit_id=audit_id, status="done", progress=1.0, report=payload
            )
        return None

    def update(self, audit_id: str, **changes: Any) -> None:
        with self._lock:
            job = self._jobs.get(audit_id)
            if job is None:
                return
            for key, value in changes.items():
                setattr(job, key, value)

    def submit(self, request: AuditRequest) -> AuditSummary:
        job = self.create()
        self._executor.submit(self._run, job.audit_id, request)
        return job

    def _run(self, audit_id: str, request: AuditRequest) -> None:
        settings = load_settings()
        cache = ResponseCache(settings.cache_path)
        client = CourtListenerClient(
            token=settings.courtlistener_token or None,
            base_url=settings.courtlistener_base_url,
            min_interval=settings.min_interval,
            max_retries=settings.max_retries,
            timeout=settings.request_timeout,
            cache=cache,
        )
        llm = (
            LLMClient(
                api_key=settings.llm_api_key,
                base_url=settings.llm_base_url,
                model=settings.llm_model,
            )
            if settings.has_llm
            else None
        )

        def progress(message: str, fraction: float) -> None:
            self.update(
                audit_id,
                status="checking" if fraction > 0.7 else "verifying",
                progress=round(fraction, 3),
                message=message,
            )

        try:
            verifier = CitationVerifier(
                client,
                llm=llm,
                options=VerifierOptions(
                    check_quotes=True,
                    check_support=request.deep and settings.has_llm,
                    progress=progress,
                ),
            )
            report = verifier.audit(request.text or "", title=request.document_title)
        except Exception as exc:  # surfaced to the client rather than swallowed
            self.update(audit_id, status="failed", error=f"{type(exc).__name__}: {exc}")
            return
        finally:
            client.close()
            cache.close()

        AUDITS_DIR.mkdir(parents=True, exist_ok=True)
        (AUDITS_DIR / f"{audit_id}.json").write_text(
            json.dumps(to_json(report), indent=2, ensure_ascii=False), encoding="utf-8"
        )
        self.update(audit_id, status="done", progress=1.0, message="Done", report=report)


jobs = AuditJobs()
app = FastAPI(title="CiteProof", version=__version__)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def _read_document_bytes(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        try:
            import docx  # provided by python-docx
        except ImportError as exc:  # pragma: no cover
            raise HTTPException(400, "docx support is not installed") from exc
        document = docx.Document(io.BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    try:
        return payload.decode("utf-8")
    except UnicodeDecodeError:
        return payload.decode("latin-1", errors="replace")


@app.get("/api/health")
def health() -> dict[str, Any]:
    settings = load_settings()
    return {
        "version": __version__,
        "case_law_access": settings.access_mode,
        "has_token": settings.has_token,
        "language_model": settings.llm_model if settings.has_llm else None,
        "support_check_available": settings.has_llm,
        "cache_path": str(settings.cache_path),
        "time": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


@app.get("/api/examples")
def examples() -> list[dict[str, str]]:
    """Demonstration documents that ship with the application."""

    items: list[dict[str, str]] = []
    if EXAMPLES_DIR.is_dir():
        for path in sorted(EXAMPLES_DIR.glob("*.txt")):
            items.append(
                {
                    "id": path.stem,
                    "title": path.stem.replace("_", " ").replace("-", " ").strip(),
                    "text": path.read_text(encoding="utf-8"),
                }
            )
    return items


@app.post("/api/extract")
def extract_only(request: AuditRequest) -> dict[str, Any]:
    """Extract citations without contacting the corpus. Useful for a first look."""

    citations, notes = extract_citations(request.text or "")
    cases = [c for c in citations if c.kind.value == "case"]
    return {
        "citations_total": len(citations),
        "case_citations": len(cases),
        "citations": [c.model_dump(mode="json") for c in citations],
        "notes": notes,
    }


@app.post("/api/audits", response_model=AuditSummary)
def create_audit(request: AuditRequest) -> AuditSummary:
    text = (request.text or "").strip()
    if not text:
        raise HTTPException(400, "The request contains no document text.")
    if len(text) > MAX_DOCUMENT_CHARS:
        raise HTTPException(413, f"The document exceeds {MAX_DOCUMENT_CHARS} characters.")
    return jobs.submit(request)


@app.post("/api/audits/file", response_model=AuditSummary)
async def create_audit_from_file(file: UploadFile = File(...)) -> AuditSummary:
    payload = await file.read()
    if not payload:
        raise HTTPException(400, "The uploaded file is empty.")
    text = _read_document_bytes(file.filename or "document.txt", payload)
    if not text.strip():
        raise HTTPException(400, "No text could be read from the uploaded file.")
    if len(text) > MAX_DOCUMENT_CHARS:
        raise HTTPException(413, f"The document exceeds {MAX_DOCUMENT_CHARS} characters.")
    title = Path(file.filename or "document").stem.replace("_", " ").replace("-", " ")
    return jobs.submit(AuditRequest(text=text, document_title=title))


@app.get("/api/audits/{audit_id}", response_model=AuditSummary)
def get_audit(audit_id: str) -> AuditSummary:
    job = jobs.get(audit_id)
    if job is None:
        raise HTTPException(404, "No audit with that id.")
    return job


@app.get("/api/audits/{audit_id}/markdown", response_class=PlainTextResponse)
def get_audit_markdown(audit_id: str) -> PlainTextResponse:
    job = jobs.get(audit_id)
    if job is None or job.report is None:
        raise HTTPException(404, "That audit has no report yet.")
    return PlainTextResponse(
        to_markdown(job.report),
        headers={"Content-Disposition": f'attachment; filename="citeproof-{audit_id}.md"'},
    )


@app.get("/api/audits/{audit_id}/json")
def get_audit_json(audit_id: str) -> JSONResponse:
    job = jobs.get(audit_id)
    if job is None or job.report is None:
        raise HTTPException(404, "That audit has no report yet.")
    return JSONResponse(
        to_json(job.report),
        headers={"Content-Disposition": f'attachment; filename="citeproof-{audit_id}.json"'},
    )


@app.get("/api/benchmark")
def benchmark() -> dict[str, Any]:
    """Measured detection results on the bundled benchmark, if it has been run."""

    if not BENCHMARK_RESULTS.is_file():
        return {
            "available": False,
            "detail": "The benchmark has not been run. Run scripts/run_benchmark.py to produce it.",
        }
    payload = json.loads(BENCHMARK_RESULTS.read_text(encoding="utf-8"))
    payload["available"] = True
    return payload


@app.get("/api/stats")
def stats() -> dict[str, Any]:
    settings = load_settings()
    cache = ResponseCache(settings.cache_path)
    count = cache.count()
    cache.close()
    return {"cached_responses": count, "cache_path": str(settings.cache_path)}


def mount_frontend(application: FastAPI) -> None:
    """Serve the built single-page application when it has been built."""

    dist = REPO_ROOT / "frontend" / "dist"
    index = dist / "index.html"
    if not index.is_file():
        return
    application.mount("/assets", StaticFiles(directory=dist / "assets"), name="assets")

    @application.get("/", response_class=HTMLResponse)
    def index_page() -> HTMLResponse:
        return HTMLResponse(index.read_text(encoding="utf-8"))

    @application.get("/favicon.svg")
    def favicon() -> FileResponse:
        return FileResponse(dist / "favicon.svg")


mount_frontend(app)
